import streamlit as st
from deepseek_api import (
    extract_knowledge,
    generate_graph_edges,
    generate_node_content,
    get_last_error,
)
from graph_utils import visualize_graph
from prompt_config import (
    DEFAULT_PROMPT_TEMPLATES,
    render_text_prompt,
)
from user_store import (
    ensure_user,
    get_user_prompt_templates,
    list_saved_graphs,
    list_users,
    save_knowledge_graph,
    save_user_prompt_templates,
)

st.set_page_config(page_title="知识图谱自动生成演示平台", layout="wide")
st.title("📚 知识图谱自动生成演示平台")

GENERATION_STATE_KEYS = (
    "edited_nodes",
    "graph_nodes",
    "graph_edges",
    "node_editor_runtime",
    "source_key",
)

def _clear_generation_state():
    for key in GENERATION_STATE_KEYS:
        st.session_state.pop(key, None)

def _as_records(rows):
    if rows is None:
        return []
    if hasattr(rows, "to_dict"):
        return rows.to_dict("records")
    return list(rows)

def _validate_nodes(rows):
    nodes = []
    errors = []
    seen_names = set()

    for index, row in enumerate(_as_records(rows), start=1):
        if not isinstance(row, dict):
            errors.append(f"第 {index} 行不是有效的知识点记录")
            continue

        name_value = row.get("name", "")
        definition_value = row.get("definition", "")
        name = "" if name_value is None else str(name_value).strip()
        definition = "" if definition_value is None else str(definition_value).strip()

        if not name:
            errors.append(f"第 {index} 行缺少 name")
            continue
        if name in seen_names:
            errors.append(f"知识点名称重复：{name}")
            continue

        seen_names.add(name)
        nodes.append({"name": name, "definition": definition})

    return nodes, errors

def _edge_to_row(edge):
    if isinstance(edge, dict):
        return {
            "起点": edge.get("source", ""),
            "终点": edge.get("target", ""),
            "关系类型": edge.get("relation_type", ""),
            "关联原因": edge.get("reason", ""),
        }
    if not isinstance(edge, (list, tuple)):
        return {"起点": "", "终点": "", "关系类型": "", "关联原因": ""}

    source = edge[0] if len(edge) >= 1 else ""
    target = edge[1] if len(edge) >= 2 else ""
    reason = edge[2] if len(edge) >= 3 else ""
    relation_type = edge[3] if len(edge) >= 4 else ""
    return {"起点": source, "终点": target, "关系类型": relation_type, "关联原因": reason}

def _decode_text_file(file_bytes):
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_bytes.decode(encoding), None
        except UnicodeDecodeError:
            continue
    return "", "文本文件编码无法识别，请使用 UTF-8、UTF-8 BOM 或 GB18030/GBK 编码。"

def _read_uploaded_text(uploaded_file):
    try:
        try:
            uploaded_file.seek(0)
        except (AttributeError, OSError):
            pass

        is_txt = uploaded_file.type == "text/plain" or uploaded_file.name.lower().endswith(".txt")
        if is_txt:
            return _decode_text_file(uploaded_file.read())

        is_pdf = uploaded_file.type == "application/pdf" or uploaded_file.name.lower().endswith(".pdf")
        if is_pdf:
            try:
                from pypdf import PdfReader
            except ImportError:
                return "", "缺少 PDF 解析依赖，请先安装 pypdf。"

            reader = PdfReader(uploaded_file)
            if reader.is_encrypted:
                try:
                    decrypt_result = reader.decrypt("")
                except Exception:
                    return "", "PDF 文件已加密，无法解析文本内容。"
                if not decrypt_result:
                    return "", "PDF 文件已加密，无法解析文本内容。"

            page_texts = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    page_texts.append(page_text.strip())

            text = "\n\n".join(page_texts)
            if not text.strip():
                return "", "未能从 PDF 中提取文本。请确认 PDF 包含可复制的文本层，扫描版 PDF 暂不支持。"
            return text, None

        from docx import Document
        doc = Document(uploaded_file)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return text, None
    except Exception as e:
        return "", f"文件解析失败：{e}"

def _validate_prompt_templates(templates):
    errors = []
    required_placeholders = {
        "knowledge_extraction": ["{{text}}"],
        "graph_edges": ["{{name_list}}", "{{text}}"],
        "node_content": ["{{name}}", "{{definition}}"],
    }

    for key in DEFAULT_PROMPT_TEMPLATES:
        value = templates.get(key, "")
        if not isinstance(value, str) or not value.strip():
            errors.append(f"{key} 不能为空")

    for key, placeholders in required_placeholders.items():
        template = templates.get(key, "")
        for placeholder in placeholders:
            if placeholder not in template:
                errors.append(f"{key} 缺少占位符 {placeholder}")

    return errors

def _select_current_user():
    st.sidebar.markdown("### 用户")
    users = list_users()
    existing_names = [user["username"] for user in users]

    if existing_names:
        selected_user = st.sidebar.selectbox("选择已有用户", existing_names + ["新用户"])
        if selected_user == "新用户":
            username = st.sidebar.text_input("新用户名", value="")
        else:
            username = selected_user
    else:
        username = st.sidebar.text_input("用户名", value="demo")

    username = username.strip()
    if not username:
        st.sidebar.warning("请输入用户名")
        st.stop()

    user = ensure_user(username)
    previous_user_id = st.session_state.get("current_user_id")
    if previous_user_id and previous_user_id != user["id"]:
        _clear_generation_state()
    st.session_state["current_user_id"] = user["id"]
    st.sidebar.caption(f"当前用户：{user['username']}")
    return user

def _render_prompt_template_page(user):
    st.markdown("### Prompt 模板配置")
    st.caption(f"当前用户：{user['username']}。模板会保存到该用户自己的配置中。")

    templates = get_user_prompt_templates(user["id"])

    with st.form("prompt_template_form"):
        system_prompt = st.text_area(
            "系统提示词：system_prompt",
            value=templates["system_prompt"],
            height=120,
        )
        knowledge_extraction = st.text_area(
            "知识点抽取模板：knowledge_extraction（占位符：{{text}}）",
            value=templates["knowledge_extraction"],
            height=240,
        )
        graph_edges = st.text_area(
            "知识关系抽取模板：graph_edges（占位符：{{name_list}}, {{text}}）",
            value=templates["graph_edges"],
            height=340,
        )
        node_content = st.text_area(
            "节点讲义模板：node_content（占位符：{{name}}, {{definition}}）",
            value=templates["node_content"],
            height=260,
        )

        col_save, col_reset = st.columns([1, 1])
        save_clicked = col_save.form_submit_button("保存模板", use_container_width=True)
        reset_clicked = col_reset.form_submit_button("恢复默认模板", use_container_width=True)

    if reset_clicked:
        save_user_prompt_templates(user["id"], DEFAULT_PROMPT_TEMPLATES)
        st.success("已恢复默认模板。")
        st.rerun()

    if save_clicked:
        next_templates = {
            "system_prompt": system_prompt,
            "knowledge_extraction": knowledge_extraction,
            "graph_edges": graph_edges,
            "node_content": node_content,
        }
        errors = _validate_prompt_templates(next_templates)
        if errors:
            st.error("模板配置存在问题：\n" + "\n".join(errors))
        else:
            save_user_prompt_templates(user["id"], next_templates)
            st.success("模板已保存到当前用户配置。")

def _render_history_page(user):
    st.markdown("### 历史知识图谱")
    st.caption(f"当前用户：{user['username']}。这里展示该用户保存过的知识点和知识图谱。")

    records = list_saved_graphs(user["id"])
    if not records:
        st.info("当前用户还没有保存过知识图谱。")
        return

    labels = [
        f"{record.get('created_at', '')} · {record.get('title', '未命名知识图谱')}"
        for record in records
    ]
    selected_index = st.selectbox("选择历史记录", range(len(records)), format_func=lambda index: labels[index])
    record = records[selected_index]
    nodes = record.get("nodes", [])
    edges = record.get("edges", [])

    st.markdown(f"**标题：** {record.get('title', '未命名知识图谱')}")
    st.markdown(f"**来源文件：** {record.get('source_filename', '未知')}")
    st.markdown(f"**保存时间：** {record.get('created_at', '')}")

    if record.get("source_excerpt"):
        with st.expander("查看原文摘录"):
            st.text(record["source_excerpt"])

    st.subheader("知识点")
    st.dataframe(nodes, use_container_width=True, hide_index=True)

    st.subheader("知识图谱")
    visualize_graph(nodes, edges)

    edge_rows = [_edge_to_row(edge) for edge in edges]
    if edge_rows:
        st.subheader("知识点关联原因")
        st.dataframe(edge_rows, use_container_width=True, hide_index=True)

    if st.button("载入到课件生成页面"):
        st.session_state["edited_nodes"] = nodes
        st.session_state["graph_nodes"] = nodes
        st.session_state["graph_edges"] = edges
        st.success("已载入到当前会话。切换到“课件生成”页面即可继续查看或生成讲义。")

def _render_graph_result(user, raw_text="", source_filename=""):
    st.subheader(" 当前知识图谱（持续可见）")
    visualize_graph(st.session_state["graph_nodes"], st.session_state["graph_edges"])
    edge_rows = [_edge_to_row(edge) for edge in st.session_state["graph_edges"]]
    if edge_rows:
        st.subheader("🔗 知识点关联原因")
        st.dataframe(edge_rows, use_container_width=True, hide_index=True)

    if raw_text and source_filename:
        save_title = st.text_input("保存标题", value=f"{source_filename} 知识图谱")
        if st.button("💾 确认保存本次知识点和知识图谱"):
            record = save_knowledge_graph(
                user["id"],
                save_title,
                source_filename,
                st.session_state["graph_nodes"],
                st.session_state["graph_edges"],
                raw_text,
            )
            st.success(f"已保存：{record['title']}")

def _render_node_content_tools(user):
    if "edited_nodes" not in st.session_state:
        return

    st.markdown("---")
    st.subheader("📚 按需生成节点讲义内容")
    content_nodes, content_node_errors = _validate_nodes(st.session_state["edited_nodes"])
    if content_node_errors:
        st.warning("知识点表格存在待修正项，讲义生成将跳过无效行：\n" + "\n".join(content_node_errors))

    for node in content_nodes:
        if st.button(f"📄 为「{node['name']}」生成内容"):
            user_templates = get_user_prompt_templates(user["id"])
            result = generate_node_content(
                node["name"],
                node.get("definition", ""),
                template=user_templates["node_content"],
                system_prompt=user_templates["system_prompt"],
            )
            if result:
                st.markdown(result)
            else:
                st.error(get_last_error() or "未能生成讲义内容，请检查 API 配置或稍后重试。")

current_user = _select_current_user()
page = st.sidebar.radio("页面", ["课件生成", "Prompt 模板配置", "历史记录"])

if page == "Prompt 模板配置":
    _render_prompt_template_page(current_user)
    st.stop()

if page == "历史记录":
    _render_history_page(current_user)
    st.stop()

st.markdown("### 步骤 1：上传教学文档")
uploaded_file = st.file_uploader("上传教学文档（TXT、DOCX 或 PDF）", type=["txt", "docx", "pdf"])

if not uploaded_file:
    if "graph_edges" in st.session_state and "graph_nodes" in st.session_state:
        st.info("当前显示的是已载入的历史知识图谱。上传新文档可重新解析。")
        _render_graph_result(current_user)
        _render_node_content_tools(current_user)
    else:
        st.info("请上传教学文档以启动内容解析与课件生成流程。")
    st.stop()

if uploaded_file:
    source_key = f"{current_user['id']}:{uploaded_file.name}:{getattr(uploaded_file, 'size', 0)}"
    if st.session_state.get("source_key") != source_key:
        _clear_generation_state()
        st.session_state["source_key"] = source_key

    raw_text, parse_error = _read_uploaded_text(uploaded_file)
    if parse_error:
        st.error(parse_error)
        st.stop()
    if not raw_text.strip():
        st.error("文件内容为空，请上传包含正文的 TXT、DOCX 或 PDF 文件。")
        st.stop()

    st.subheader("📖 文本预览")
    st.text_area("教学内容", raw_text, height=200)

    if st.button("🔍 解析并提取知识点"):
        with st.spinner("调用 DeepSeek 模型中..."):
            user_templates = get_user_prompt_templates(current_user["id"])
            full_prompt = render_text_prompt(user_templates["knowledge_extraction"], raw_text)
            nodes = extract_knowledge(full_prompt, system_prompt=user_templates["system_prompt"])
            nodes, node_errors = _validate_nodes(nodes)
            if not nodes:
                error_detail = get_last_error()
                if error_detail:
                    st.error(f"未能成功提取知识点：{error_detail}")
                else:
                    st.error("未能成功提取知识点，请检查 Prompt 或内容格式。")
            else:
                if node_errors:
                    st.warning("部分知识点格式不完整，已自动跳过：\n" + "\n".join(node_errors))
                st.success("成功提取知识点，可在下方编辑")
                st.session_state["edited_nodes"] = nodes

    if "edited_nodes" in st.session_state:
        st.subheader("📋 当前知识点（可再次编辑）")
        st.session_state["edited_nodes"] = st.data_editor(
            st.session_state["edited_nodes"],
            key="node_editor_runtime",
        )
        st.markdown("---")
        st.subheader("📡 点击生成知识图谱")
        if st.button("🧠 生成知识图谱"):
            with st.spinner("生成图谱中..."):
                graph_nodes, node_errors = _validate_nodes(st.session_state["edited_nodes"])
                if node_errors:
                    st.error("请先修正知识点表格：\n" + "\n".join(node_errors))
                else:
                    user_templates = get_user_prompt_templates(current_user["id"])
                    edges = generate_graph_edges(
                        raw_text,
                        [n["name"] for n in graph_nodes],
                        graph_edges_template=user_templates["graph_edges"],
                        system_prompt=user_templates["system_prompt"],
                    )
                    if not edges:
                        error_detail = get_last_error()
                        if error_detail:
                            st.warning(f"未能生成知识关系：{error_detail}")
                        else:
                            st.warning("未生成知识关系，可先保存知识点，或调整 Prompt 后重新生成。")
                    st.session_state["edited_nodes"] = graph_nodes
                    st.session_state["graph_edges"] = edges
                    st.session_state["graph_nodes"] = graph_nodes

    if "graph_edges" in st.session_state:
        _render_graph_result(current_user, raw_text, uploaded_file.name)

    if "edited_nodes" in st.session_state:
        _render_node_content_tools(current_user)
